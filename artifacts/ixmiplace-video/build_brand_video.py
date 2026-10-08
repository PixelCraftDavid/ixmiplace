import build_template as b
from pptx import Presentation
from pptx.util import Inches
from PIL import Image
from pathlib import Path

b.R=Presentation(); b.R.slide_width=Inches(6); b.R.slide_height=Inches(10.666667)
b.BG='F4F0E6'; b.INK='1C211A'; b.CYAN='E4B66E'; b.PURPLE='477450'; b.MUTED='78746A'
O=Path(__file__).parent
T=Path('C:/Users/DAVIDP~1/AppData/Local/Temp')
files={
 'home':'codex-clipboard-751b6647-d926-449b-8269-0a70ac4b94d2.png',
 'listings':'codex-clipboard-d1ac2ad3-768f-472e-8831-24d354bf4697.png',
 'map':'codex-clipboard-20eb3f1e-2b73-45ce-96d4-257bbac76613.png',
 'mobile':'codex-clipboard-6b554ee0-a0cb-4165-9b6b-89e4b1fa04fa.png',
 'contact':'codex-clipboard-67df62bc-b9d1-431c-b4aa-69d4ed20a261.png',
 'visit':'codex-clipboard-b4b49a80-7168-4b6e-8fcf-8ba837d67ac5.png'}

def pic(s,key,x,y,w,h,crop=None):
    path=T/files[key]; iw,ih=Image.open(path).size
    q=s.shapes.add_picture(str(path),Inches(x),Inches(y),width=Inches(w),height=Inches(h))
    if crop:
        l,t,r,bot=crop
        q.crop_left=l/iw; q.crop_top=t/ih; q.crop_right=(iw-r)/iw; q.crop_bottom=(ih-bot)/ih
    return q

def laptop(s,x,y,w,h,key='home',crop=None):
    b.shape(s,x-.04,y+.1,w+.08,h+.05,'D5DCCF')
    b.shape(s,x,y,w,h,b.INK)
    pic(s,key,x+.1,y+.13,w-.2,h-.27,crop)
    b.shape(s,x-.2,y+h-.02,w+.4,.15,'BBC8B1')

def mobile(s,x,y,w=2.55,h=5.08,key='mobile',crop=None):
    b.shape(s,x-.06,y+.1,w+.12,h+.08,'D5DCCF')
    b.shape(s,x,y,w,h,b.INK)
    pic(s,key,x+.09,y+.1,w-.18,h-.2,crop)

def hero_phone(s,x,y,w=2.7,h=5.0):
    b.shape(s,x-.06,y+.1,w+.12,h+.07,'D5DCCF')
    b.shape(s,x,y,w,h,b.INK)
    pic(s,'home',x+.08,y+.09,w-.16,h-.18,(930,61,1907,564))
    b.shape(s,x+.11,y+.12,w-.22,.57,b.BG)
    b.text(s,'IxmiPlace',x+.24,y+.27,w-.5,.25,13,b.INK,True,font='League Spartan')
    b.shape(s,x+.11,y+.74,w-.22,2.01,'1C211A')
    b.text(s,'Encuentra tu\npróximo espacio en\nIxmiquilpan.',x+.27,y+1,w-.5,1.2,23,b.WHITE,True,font='League Spartan')
    b.text(s,'Casas, rentas y hospedaje\nsin intermediarios.',x+.27,y+2.17,w-.5,.52,9,b.WHITE)
    b.shape(s,x+.19,y+3.08,w-.38,1.1,b.BG)
    b.text(s,'¿Qué buscas?\nZona · Ixmiquilpan',x+.35,y+3.26,w-.68,.55,11,b.INK,True)
    b.pill(s,'Buscar',x+.23,y+h-.65,w-.46,'477450',b.WHITE,11)

s,tc=b.base(1,False,5)
b.text(s,'¿Buscando casa\no roomie?',.42,1.02,5.2,1.58,38,tc,True,font='League Spartan')
b.pill(s,'ENCUENTRA TU ESPACIO EN IXMIQUILPAN',.42,2.7,4.75,'477450',b.WHITE,10)
hero_phone(s,.59,3.57,2.7,5.55)
laptop(s,3.19,4.67,2.32,1.46,'home',(0,61,1912,565))
b.shape(s,3.49,6.58,1.78,1.65,'26392F')
b.text(s,'Tu comunidad.\nTu hogar.',3.7,6.97,1.38,1.05,17,'E4B66E',True,font='League Spartan')
b.pill(s,'CLARO ↔ OSCURO',3.4,8.62,2.12,'7EAA9B',b.INK,9)

s,tc=b.base(2,True,7)
b.text(s,'Filtros reales.\nMapas exactos.',.42,1.04,5.2,1.6,37,tc,True,font='League Spartan')
b.pill(s,'RENTA',.42,2.76,1.3,'E4B66E',b.INK,11)
b.pill(s,'VENTA',1.88,2.76,1.3,'7EAA9B',b.INK,11)
laptop(s,.52,3.52,4.96,2.19,'map',(38,170,1874,841))
mobile(s,3.47,6.11,1.91,3.17,'listings',(39,87,484,527))
b.text(s,'Viviendas,\nlocales y roomies\nen un solo lugar.',.45,6.43,2.82,1.42,22,b.WHITE,True,font='League Spartan')
b.pill(s,'EXPLORA EL MAPA',.45,8.62,2.77,'477450',b.WHITE,10)

s,tc=b.base(3,False,8)
b.text(s,'Contacto directo\npor WhatsApp',.42,1.04,5.2,1.58,34,tc,True,font='League Spartan')
b.text(s,'+ Guías imprimibles.',.42,2.65,5.2,.63,24,'477450',True,font='League Spartan')
b.shape(s,.42,3.62,5.16,1.34,b.WHITE)
pic(s,'contact',.58,3.74,4.84,1.04,(1121,138,1401,283))
b.shape(s,.49,5.44,2.39,3.4,b.WHITE,line='DDE4D8')
b.text(s,'GUÍA IXMIPLACE',.68,5.66,1.98,.26,10,'477450',True)
b.text(s,'Lista para revisar\nantes de decidir',.68,6.15,1.99,.86,18,b.INK,True,font='League Spartan')
for j,t in enumerate(['Costos y requisitos','Estado del inmueble','Servicios incluidos','Visita y ubicación']):
    b.shape(s,.7,7.25+j*.28,.11,.11,'E4B66E',b.MSO_SHAPE.RECTANGLE)
    b.text(s,t,.94,7.21+j*.28,1.75,.22,8,b.INK)
b.shape(s,3.15,5.75,2.38,2.68,b.WHITE,line='DDE4D8')
pic(s,'visit',3.3,5.91,2.08,1.51,(26,118,492,348))
b.pill(s,'COMPROBANTE DE VISITA',3.3,7.87,2.08,'EFF2EC','477450',8)
b.pill(s,'Español',.42,9.17,1.46,'477450',b.WHITE,10)
b.pill(s,'English',2.03,9.17,1.46,'7EAA9B',b.INK,10)
b.pill(s,'Hñähñu',3.64,9.17,1.62,'E4B66E',b.INK,10)

s,tc=b.base(4,True,7)
b.text(s,'Promociona tu\nnegocio local',.42,1.04,5.2,1.6,36,tc,True,font='League Spartan')
b.text(s,'desde $99/mes.',.42,2.68,5.2,.69,30,'E4B66E',True,font='League Spartan')
mobile(s,.51,3.64,2.5,4.84,'mobile',(12,57,367,520))
for i,(price,c) in enumerate([('$99','E4B66E'),('$199','7EAA9B'),('$349','F4F0E6')]):
    y=3.83+i*1.57
    b.shape(s,3.4,y,2.09,1.32,c)
    b.text(s,price,3.57,y+.26,1.75,.53,31,b.INK,True,'center',font='League Spartan')
    b.text(s,'MXN / MES',3.57,y+.91,1.75,.24,9,b.INK,True,'center')
b.text(s,'Métricas en tiempo real\ny soporte directo.',.45,8.81,5.1,.85,21,b.WHITE,True,font='League Spartan')

s,tc=b.base(5,False,8)
b.text(s,'¡Todo en\nIxmiPlace!',.42,1.04,5.2,1.77,43,tc,True,font='League Spartan')
laptop(s,.54,3.42,4.91,2.12,'home',(0,61,1912,565))
mobile(s,3.41,5.45,2.02,3.21,'mobile',(12,57,367,520))
s.shapes.add_picture(str(O.parent.parent/'public'/'logo-ixmiplace.jpg'), Inches(.71), Inches(6.08), width=Inches(2.07),height=Inches(2.12))
b.text(s,'Entra ya:',.44,8.92,1.2,.34,13,b.MUTED,True)
b.text(s,'ixmiplace.vercel.app',1.55,8.9,4.12,.47,20,b.INK,True,font='League Spartan')
b.shape(s,.44,9.44,5.12,.47,'477450')
b.text(s,'EXPLORAR AHORA  →',.63,9.56,4.74,.25,12,b.WHITE,True,'center')
b.R.core_properties.title='IxmiPlace | Video 35s | Identidad original'
b.R.save(O/'IxmiPlace-Video-35s-Colores-Reales.pptx')
print(O/'IxmiPlace-Video-35s-Colores-Reales.pptx')

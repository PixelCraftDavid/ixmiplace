from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from pathlib import Path

OUT = Path(__file__).parent
R = Presentation()
R.slide_width = Inches(6)
R.slide_height = Inches(10.666667)
BG='F7F9FC'; INK='101B33'; CYAN='24CFE0'; PURPLE='7856FF'; WHITE='FFFFFF'; MUTED='768397'

def shape(s,x,y,w,h,c,kind=MSO_SHAPE.ROUNDED_RECTANGLE,line=None):
    z=s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    z.fill.solid(); z.fill.fore_color.rgb=RGBColor.from_string(c)
    z.line.fill.background() if line is None else None
    if line: z.line.color.rgb=RGBColor.from_string(line); z.line.width=Pt(1)
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        z.adjustments[0]=0.13
    return z

def text(s,txt,x,y,w,h,size=18,c=INK,bold=False,align='left',font='Inter'):
    z=s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    f=z.text_frame; f.clear(); f.word_wrap=True
    f.margin_left=f.margin_right=0; f.margin_top=f.margin_bottom=0
    for i,line in enumerate(txt.split('\n')):
        p=f.paragraphs[0] if i==0 else f.add_paragraph()
        p.text=line; p.font.name=font; p.font.size=Pt(size); p.font.bold=bold; p.font.color.rgb=RGBColor.from_string(c)
        p.alignment={'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}[align]
        p.space_after=Pt(0); p.space_before=Pt(0)
    return z

def line(s,x,y,x2,y2,c,width=2):
    from pptx.enum.shapes import MSO_CONNECTOR
    q=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y), Inches(x2), Inches(y2))
    q.line.color.rgb=RGBColor.from_string(c); q.line.width=Pt(width)

def pill(s,label,x,y,w,c=CYAN,tc=INK,size=12):
    shape(s,x,y,w,.38,c)
    text(s,label,x+.08,y+.09,w-.16,.22,size,tc,True,'center')

def base(n,dark,duration):
    s=R.slides.add_slide(R.slide_layouts[6])
    s.background.fill.solid(); s.background.fill.fore_color.rgb=RGBColor.from_string(INK if dark else BG)
    tc=WHITE if dark else INK
    shape(s,.38,.4,.12,.32,CYAN)
    text(s,'IXMIPLACE',.63,.43,3,.25,11,tc,True,font='League Spartan')
    text(s,f'0{n} / 05',4.7,.44,.9,.22,9,MUTED,align='right')
    shape(s,.38,9.97,5.24,.02,'33415A' if dark else 'DEE5EF',MSO_SHAPE.RECTANGLE)
    text(s,'TU COMUNIDAD, TU HOGAR',.4,10.12,3.4,.24,9,MUTED,True)
    text(s,'IXMIQUILPAN',4.15,10.12,1.45,.24,9,MUTED,align='right')
    t=OxmlElement('p:transition'); t.set('advClick','0'); t.set('advTm',str(duration*1000)); t.append(OxmlElement('p:cut'))
    s._element.append(t)
    s.notes_slide.notes_text_frame.text=f'Escena {n}. Duración {duration}s. Animar título con Pop. Gráficos originales editables, sin recursos premium. Agregar música electrónica instrumental gratuita verificada para uso comercial.'
    return s,tc

def house(s,x,y,w,h,c=WHITE):
    shape(s,x,y+h*.32,w,h*.68,c,MSO_SHAPE.RECTANGLE)
    shape(s,x-w*.05,y,w*1.1,h*.5,c,MSO_SHAPE.ISOSCELES_TRIANGLE)
    shape(s,x+w*.43,y+h*.61,w*.22,h*.39,INK,MSO_SHAPE.RECTANGLE)
    shape(s,x+w*.12,y+h*.48,w*.17,h*.18,CYAN,MSO_SHAPE.RECTANGLE)

def pin(s,x,y,c=PURPLE,scale=1):
    shape(s,x,y,.22*scale,.22*scale,c,MSO_SHAPE.OVAL)
    q=shape(s,x+.03*scale,y+.12*scale,.16*scale,.18*scale,c,MSO_SHAPE.ISOSCELES_TRIANGLE); q.rotation=180
    shape(s,x+.07*scale,y+.07*scale,.08*scale,.08*scale,WHITE,MSO_SHAPE.OVAL)

def map_ui(s,x,y,w,h,dark=False):
    shape(s,x,y,w,h,'25364A' if dark else 'E4EFEE')
    # schematic river and streets, all editable vectors
    line(s,x+w*.58,y+.08,x+w*.4,y+h-.08,'53CFE0',16)
    for a in [.24,.55,.78]: line(s,x+.08,y+h*a,x+w-.08,y+h*a,'42556B' if dark else WHITE,7)
    for a in [.18,.72]: line(s,x+w*a,y+.08,x+w*a,y+h-.08,'42556B' if dark else WHITE,7)
    for a,b in [(.17,.23),(.66,.36),(.31,.64),(.75,.77),(.43,.4)]: pin(s,x+w*a,y+h*b)
    pill(s,'Ixmiquilpan',x+.13,y+h-.46,1.1,WHITE,INK,8)

def phone(s,x,y,w=2.35,h=4.7,dark=False,mode='listing'):
    shape(s,x-.08,y+.12,w+.16,h+.09,'DCE2EB')
    shape(s,x,y,w,h,'111A2A')
    shape(s,x+.09,y+.1,w-.18,h-.2,'17243A' if dark else WHITE)
    shape(s,x+w*.35,y+.14,w*.3,.09,'111A2A')
    tc=WHITE if dark else INK
    text(s,'IxmiPlace',x+.2,y+.37,w-.4,.28,14,tc,True,font='League Spartan')
    pill(s,'Buscar por zona',x+.18,y+.79,w-.36,'2B3A51' if dark else 'EDF1F7',tc,9)
    if mode=='map':
        pill(s,'Renta',x+.18,y+1.28,.72,CYAN,INK,8)
        pill(s,'Venta',x+1,y+1.28,.72,'33435B' if dark else 'EDF1F7',tc,8)
        map_ui(s,x+.17,y+1.86,w-.34,h-2.35,dark)
    else:
        shape(s,x+.18,y+1.34,w-.36,1.15,PURPLE)
        house(s,x+.77,y+1.55,.67,.72)
        text(s,'Encuentra tu espacio',x+.2,y+2.68,w-.4,.4,11,tc,True)
        text(s,'Viviendas · Locales · Roomies',x+.2,y+3.12,w-.4,.4,8,tc)
        pill(s,'Explorar ahora',x+.18,y+h-.75,w-.36,CYAN,INK,10)
    return x,y,w,h

def laptop(s,x,y,w=4.7,h=2.9,mode='map'):
    shape(s,x,y,w,h,'253248')
    shape(s,x+.1,y+.12,w-.2,h-.28,WHITE)
    text(s,'IxmiPlace',x+.25,y+.23,w-.5,.25,12,INK,True,font='League Spartan')
    if mode=='map':
        pill(s,'Renta',x+.22,y+.7,.8,CYAN,INK,8)
        pill(s,'Venta',x+1.12,y+.7,.8,'EDF1F7',INK,8)
        map_ui(s,x+.2,y+1.21,w-.4,h-1.47)
    else:
        shape(s,x+.23,y+.71,w-.46,h-1.03,PURPLE)
        text(s,'Tu comunidad,\ntu hogar.',x+.46,y+1.09,w*.5,1,23,WHITE,True,font='League Spartan')
        house(s,x+w*.72,y+1.02,.6,.8)
    shape(s,x-.2,y+h-.02,w+.4,.17,'B9C5D7')

# 1 / HOOK
s,tc=base(1,False,5)
text(s,'¿Buscando casa\no roomie?',.42,1.08,5.2,1.6,38,tc,True,font='League Spartan')
pill(s,'UN SOLO LUGAR. MÁS POSIBILIDADES.',.42,2.76,4.2,INK,WHITE,10)
phone(s,.67,3.59,2.4,5.05,False)
phone(s,3.11,4.03,2.25,4.76,True)
pill(s,'CLARO',.94,9.03,1.2,'E5EBF3',INK,9)
pill(s,'OSCURO',3.58,9.03,1.4,INK,WHITE,9)

# 2 / MAP AND FILTERS
s,tc=base(2,True,7)
text(s,'Filtros reales.\nMapas exactos.',.42,1.08,5.2,1.6,37,tc,True,font='League Spartan')
pill(s,'RENTA',.42,2.75,1.3,CYAN,INK,11)
pill(s,'VENTA',1.88,2.75,1.3,PURPLE,WHITE,11)
laptop(s,.53,3.6,4.9,3.13)
phone(s,3.31,5.4,2.02,3.9,True,'map')
shape(s,.44,7.19,2.54,1.52,'1D2C43')
text(s,'Viviendas,\nlocales y roomies\nen un solo lugar.',.64,7.44,2.17,1.13,16,WHITE,True)
pill(s,'ENCUENTRA TU ZONA',.42,9.27,3.1,CYAN,INK,10)

# 3 / DIRECT CONTACT + GUIDES
s,tc=base(3,False,8)
text(s,'Contacto directo\npor WhatsApp',.42,1.08,5.2,1.6,34,tc,True,font='League Spartan')
text(s,'+ Guías imprimibles.',.42,2.68,5.2,.6,23,PURPLE,True,font='League Spartan')
phone(s,.61,3.66,2.41,4.87,False)
shape(s,3.12,4.17,2.19,3.1,'DDE3ED')
shape(s,2.98,3.99,2.19,3.1,WHITE,line='DCE3ED')
text(s,'GUÍA LOCAL',3.2,4.23,1.8,.28,13,INK,True,font='League Spartan')
text(s,'Tu próximo hogar',3.2,4.71,1.7,.3,10,INK)
for j in range(5): shape(s,3.2,5.25+j*.23,1.56-(j%2)*.34,.045,'DCE3ED',MSO_SHAPE.RECTANGLE)
pin(s,4.31,6.37,CYAN,1.6)
shape(s,2.58,7.3,1.2,1.2,'25D366',MSO_SHAPE.OVAL)
text(s,'WA',2.76,7.64,.85,.43,23,WHITE,True,'center',font='League Spartan')
pill(s,'Español',.42,9.15,1.44,INK,WHITE,10)
pill(s,'English',2.01,9.15,1.43,PURPLE,WHITE,10)
pill(s,'Hñähñu',3.59,9.15,1.65,CYAN,INK,10)

# 4 / LOCAL BUSINESS
s,tc=base(4,True,7)
text(s,'Promociona tu\nnegocio local',.42,1.08,5.2,1.6,36,tc,True,font='League Spartan')
text(s,'desde $99/mes.',.42,2.73,5.2,.69,29,CYAN,True,font='League Spartan')
shape(s,.4,3.7,5.2,2.35,'1E2C43')
text(s,'NEGOCIOS DE TU COMUNIDAD',.65,3.99,4.7,.25,10,CYAN,True)
for x,c,txt in [(.66,PURPLE,'Tu negocio'),(2.24,'257784','Aquí'),(3.82,'3C4D69','Destaca')]:
    shape(s,x,4.51,1.46,1.19,c)
    house(s,x+.46,4.66,.5,.49)
    text(s,txt,x+.07,5.28,1.32,.22,9,WHITE,True,'center')
for i,(price,c) in enumerate([('$99',CYAN),('$199',PURPLE),('$349','FFFFFF')]):
    x=.4+i*1.79
    shape(s,x,6.52,1.62,1.62,c)
    text(s,price,x+.09,6.94,1.44,.5,27,INK if i!=1 else WHITE,True,'center',font='League Spartan')
    text(s,'MXN / MES',x+.08,7.65,1.46,.23,9,INK if i!=1 else WHITE,True,'center')
text(s,'Métricas en tiempo real\ny soporte directo.',.44,8.65,5.1,.85,21,WHITE,True,font='League Spartan')

# 5 / CLOSE
s,tc=base(5,False,8)
text(s,'¡Todo en\nIxmiPlace!',.42,1.08,5.2,1.75,43,tc,True,font='League Spartan')
laptop(s,.56,3.56,4.82,2.92,'brand')
phone(s,3.49,5.0,1.92,3.87,False)
s.shapes.add_picture(str(OUT.parent.parent/'public'/'logo-ixmiplace.jpg'), Inches(.68), Inches(6.89), width=Inches(1.6), height=Inches(1.64))
text(s,'Entra ya:',.44,8.97,1.2,.34,13,MUTED,True)
text(s,'ixmiplace.vercel.app',1.55,8.93,4.12,.47,20,INK,True,font='League Spartan')
shape(s,.44,9.48,5.12,.41,CYAN)
text(s,'EXPLORAR AHORA  →',.63,9.56,4.74,.25,12,INK,True,'center')

R.core_properties.title='IxmiPlace | Anuncio vertical 35s | Recursos gratuitos'
R.core_properties.subject='Plantilla editable de cinco escenas para terminar en Canva'
R.core_properties.author='IxmiPlace'
R.save(OUT/'IxmiPlace-Video-35s-Editable.pptx')
print(OUT/'IxmiPlace-Video-35s-Editable.pptx')
